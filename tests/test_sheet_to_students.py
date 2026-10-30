# Run: python3 -m unittest discover -s tests
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
import sheet_to_students as sts  # noqa: E402

HEADER = ["Номер тел", "ИИН ученика / TUMO ID", "Новое расписание (при измении)", "Почта TUMO", "Расписание",
          "ФИО ученика", "День SL", "ex-Coach", "Новый коуч SL", "Ушел к коучу:", "Изменен график в Т360",
          "Отправлено соообщение", "Причина ухода"]


def row(name, day, ex, coach, email="", new="", sched="", phone="87001234567", iin="130101500123"):
    return [phone, iin, new, email, sched, name, day, ex, coach, None, "TRUE", "FALSE", "note"]


def xlsx(tabs, fills=None):
    """fills: {(tab, row number): "RRGGBB"} — fill the whole row with that colour."""
    import openpyxl
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in tabs.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        for (t, line), rgb in (fills or {}).items():
            if t == title:
                for cell in ws[line]:
                    cell.fill = PatternFill("solid", fgColor="FF" + rgb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class SheetToStudentsTest(unittest.TestCase):
    def setUp(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")

    def test_maps_columns_and_drops_private_fields(self):
        raw = xlsx({
            "ПН ЧТ 1030": [HEADER,
                           row("БАТЫРЖАН НҰРАЛИ АМАНГЕЛДІҰЛЫ", "Четверг ", "gulnaz", "Аружан Т",
                               email=" Batyrzhan.N@tumo.world ", sched="Понедельник, Четверг : 10:30–12:97"),
                           row("Иван Петров", "Понедельник", "Nazym", "Almas", email="ivan@gmail.com"),
                           row("Сломанный День", "Пятница", "", "Almas"),
                           [None] * 13],
            "Цифры": [["ON", 531]],
        })
        students, problems = sts.parse(sts.read_tables(raw))
        self.assertEqual(students, [
            ["Батыржан Нұрали Амангелдіұлы", "Четверг", "10:30-12:30", "Aruzhan Torebek", "Gulnaz",
             "Самообучение", "batyrzhan.n@tumo.world", "", "", []],
            ["Иван Петров", "Понедельник", "10:30-12:30", "Almas", "Nazym", "Самообучение", "", "", "", []],
        ])
        self.assertEqual(len(problems), 1)
        self.assertIn("Сломанный День", problems[0])
        flat = repr(students)
        self.assertNotIn("8700", flat)
        self.assertNotIn("130101", flat)

    def test_duplicate_prefers_confirmed_new_schedule_then_published(self):
        tabs = {
            "ПН ЧТ 1030": [HEADER, row("Сара Мухтар", "Понедельник", "Gulnaz", "Gulnaz")],
            "ВТ ПТ 1430": [HEADER, row("Сара Мухтар", "Пятница", "Gulnaz", "Alinur",
                                       new="Вторник, Пятница : 14:30–16:30")],
        }
        students, problems = sts.parse(sts.read_tables(xlsx(tabs)))
        self.assertEqual(students[0][1:4], ["Пятница", "14:30-16:30", "Alinur"])
        self.assertEqual(len(problems), 1)

        tabs["ВТ ПТ 1430"][1][2] = ""  # no confirmation -> keep what is published
        published = {sts.norm("Сара Мухтар"): ["Сара Мухтар", "Понедельник", "10:30-12:30"]}
        students, _ = sts.parse(sts.read_tables(xlsx(tabs)), published)
        self.assertEqual(students[0][1:3], ["Понедельник", "10:30-12:30"])

    def test_csv_single_tab_takes_time_from_schedule(self):
        csv_text = ",".join(HEADER) + "\n" + ",".join(
            '"%s"' % v for v in row("ДАРХАН АСКАРОВ", "Вторник", "Nazym", "Gulnaz",
                                    sched="Вторник, Пятница : 16:30–18:30.")) + "\n"
        students, _ = sts.parse(sts.read_tables(csv_text.encode("utf-8")))
        self.assertEqual(students[0][:3], ["Дархан Аскаров", "Вторник", "16:30-18:30"])

    def test_latin_lookalikes_in_cyrillic_names_are_fixed(self):
        self.assertEqual(sts.fix_mixed_script("Tөлеуғазы Алдияр"), "Төлеуғазы Алдияр")
        self.assertEqual(sts.fix_mixed_script("Iңкар Есiмхан Mади"), "Іңкар Есімхан Мади")
        self.assertEqual(sts.fix_mixed_script("Agayeva Zakhirat"), "Agayeva Zakhirat")
        raw = xlsx({"ПН ЧТ 1630": [HEADER, row("TӨЛЕУҒАЗЫ АЛДИЯР", "Четверг", "Nurai", "Elmira")]})
        students, _ = sts.parse(sts.read_tables(raw))
        self.assertEqual(students[0][0], "Төлеуғазы Алдияр")

    def test_row_colours_give_the_status(self):
        tabs = {"ВТ ПТ 1430": [HEADER,
                               row("Активный Ученик", "Вторник", "Almas", "Almas", email="a@tumo.world"),
                               row("Красный Ученик", "Пятница", "Almas", "Almas", email="r@tumo.world"),
                               row("Светло Красный", "Пятница", "Almas", "Almas"),
                               row("Оранжевый Ученик", "Вторник", "Almas", "Almas"),
                               row("Жёлтый Ученик", "Вторник", "Almas", "Almas"),
                               row("Фиолетовый Ученик", "Вторник", "Almas", "Almas", email="p@tumo.world"),
                               row("Тёмно Фиолетовый", "Вторник", "Almas", "Almas"),
                               row("Светло Фиолетовый", "Вторник", "Almas", "Almas"),
                               row("Пурпурный Ученик", "Вторник", "Almas", "Almas")]}
        fills = {("ВТ ПТ 1430", 3): "FF0000", ("ВТ ПТ 1430", 4): "E06666",
                 ("ВТ ПТ 1430", 5): "FF9900", ("ВТ ПТ 1430", 6): "FFFF00",
                 ("ВТ ПТ 1430", 7): "9900FF", ("ВТ ПТ 1430", 8): "674EA7",
                 ("ВТ ПТ 1430", 9): "B4A7D6", ("ВТ ПТ 1430", 10): "FF00FF"}
        stats = {}
        students, _ = sts.parse(sts.read_tables(xlsx(tabs, fills)), stats=stats)
        status = {s[0]: s[8] for s in students}
        self.assertEqual(status, {"Активный Ученик": "", "Красный Ученик": "inactive", "Светло Красный": "inactive",
                                  "Оранжевый Ученик": "", "Жёлтый Ученик": "", "Фиолетовый Ученик": "waitlist",
                                  "Тёмно Фиолетовый": "waitlist", "Светло Фиолетовый": "", "Пурпурный Ученик": ""})
        # inactive / waitlist rows keep only the name and stage
        red = next(s for s in students if s[0] == "Красный Ученик")
        self.assertEqual(red, ["Красный Ученик", "", "", "", "", "Самообучение", "", "", "inactive", []])
        purple = next(s for s in students if s[0] == "Фиолетовый Ученик")
        self.assertEqual(purple, ["Фиолетовый Ученик", "", "", "", "", "Самообучение", "", "", "waitlist", []])
        self.assertEqual(stats, {"red": 2, "purple": 2, "notes_inactive": 0, "notes_waitlist": 0})

    def test_active_row_wins_over_red_or_purple_old_row(self):
        tabs = {"ПН ЧТ 1030": [HEADER, row("Сара Мухтар", "Понедельник", "Gulnaz", "Gulnaz")],
                "ВТ ПТ 1430": [HEADER, row("Сара Мухтар", "Пятница", "Gulnaz", "Alinur")],
                "ВТ ПТ 1630": [HEADER, row("Сара Мухтар", "Вторник", "Gulnaz", "Gulnaz")]}
        fills = {("ПН ЧТ 1030", 2): "FF0000", ("ВТ ПТ 1630", 2): "9900FF"}
        students, problems = sts.parse(sts.read_tables(xlsx(tabs, fills)))
        self.assertEqual(students, [["Сара Мухтар", "Пятница", "14:30-16:30", "Alinur", "Gulnaz", "Самообучение", "", "", "", []]])
        self.assertEqual(problems, [])  # an old red/purple copy is not an ambiguity
        # only red + purple rows: waitlist wins over inactive
        only_old = {k: v for k, v in tabs.items() if k != "ВТ ПТ 1430"}
        students, _ = sts.parse(sts.read_tables(xlsx(only_old, fills)))
        self.assertEqual(students[0][8], "waitlist")

    def test_inactive_keywords_in_notes_columns(self):
        header = HEADER + [None]  # an unlabelled note column after «Причина ухода»
        def r(name, reason="", extra=None, sched=""):
            cells = row(name, "Вторник", "Almas", "Almas", sched=sched)
            cells[12] = reason
            return cells + [extra]
        tabs = {"ВТ ПТ 1030": [header,
                               r("Первый Ученик", reason="ignore"),
                               r("Второй Ученик", reason="self initiated"),
                               r("Третий Ученик", reason="ЛИСТ ОЖИДАНИЯ 2027"),
                               r("Четвёртый Ученик", extra="отчислен"),
                               r("Пятый Ученик", sched="Вторник, Пятница : 10:30–12:30 (лист ожидания)"),
                               r("Шестой Ученик", reason="отказ, не успевает по учебе")]}
        stats = {}
        students, _ = sts.parse(sts.read_tables(xlsx(tabs)), stats=stats)
        # «Расписание» with "(лист ожидания)" is the originally requested slot, not a status
        self.assertEqual({s[0]: s[8] for s in students}, {
            "Первый Ученик": "inactive", "Второй Ученик": "inactive", "Третий Ученик": "waitlist",
            "Четвёртый Ученик": "inactive", "Пятый Ученик": "", "Шестой Ученик": "inactive"})
        self.assertEqual(stats, {"red": 0, "purple": 0, "notes_inactive": 4, "notes_waitlist": 1})

    def test_waitlist_notes_exclude_only_students_not_attending_now(self):
        out = ["ЛИСТ ОЖИДАНИЯ 2027", "ЛИСТ ОЖИДАНИЯ ЛЕТО 2027", "лист ожидания ноябрь", "лист ожидания - лето",
               "лист ожидания весна 2027", "Лист ожидания на ноябрь просит 16:30", "лист ожидания (16 30 в ноябре)",
               "хотят на 16:30 на след поток", "на новый поток", "хотят после НГ начать обучение",
               "не может посещать", "не сможет посещать, хочет в лист ожидания на лето", "отказ",
               "лист ожидание, отказ", "приостановлено обучени в связи с лечением", "wants to suspend till next year",
               "ignore", "self initiated", "отчислен"]
        keep = ["лист ожидания", "лист ожидания (хочет пн/чт 16:30)", "лист ожидания (хочет на 14:30)",
                "лист ожидания хочет на 16:30", "хочет на 16:30 лист ожидания", "ЛИСТ ОЖИДАНИЯ СЛ",
                "ЛИСТ ОЖИДАНИЯ 16:30", "wl 16:30 tue-fri", "хочет вт/пт 14:30", "не ответил",
                "пока будет ходить 14:30, хочет на пн/чт 16:30-18:30 (можно на сл)", "ребенок не в городе",
                "Новый коуч- отправить сообщение", "Almas Support must", "should be back in November"]
        for note in out:
            self.assertTrue(sts.INACTIVE_RE.search(note), note)
        for note in keep:
            self.assertFalse(sts.INACTIVE_RE.search(note), note)
            self.assertEqual(sts.note_status(note), "", note)
        # left -> inactive (red-row message); paused / later waitlist -> waitlist (suspended message)
        for note in ("отказ", "лист ожидание, отказ", "не может посещать", "ignore", "self initiated", "отчислен"):
            self.assertEqual(sts.note_status(note), "inactive", note)
        for note in ("ЛИСТ ОЖИДАНИЯ 2027", "лист ожидания ноябрь", "приостановлено обучени", "на новый поток",
                     "wants to suspend till next year", "хотят после НГ начать обучение"):
            self.assertEqual(sts.note_status(note), "waitlist", note)

    def test_is_red(self):
        from openpyxl.styles.colors import Color
        for rgb in ("FFFF0000", "FFE06666", "FFCC0000", "FFEA4335"):
            self.assertTrue(sts.is_red(Color(rgb=rgb)), rgb)
        for rgb in ("FFFF9900", "FFFF00FF", "FFFFFF00", "FFEA9999", "FFFFFFFF", "FF9900FF", "FF00FF00"):
            self.assertFalse(sts.is_red(Color(rgb=rgb)), rgb)
        self.assertTrue(sts.is_red(Color(indexed=10)))   # palette red
        self.assertFalse(sts.is_red(Color(theme=5)))     # theme colours aren't resolved
        self.assertFalse(sts.is_red(None))

    def test_onboarding_stream_tab(self):
        header = ["", "Номер родителя", "почта родителя", "ФИО ребенка", "Номер ребенка", "ДР ребенка",
                  "Расписание", "Коуч", "", "Зона обучения", "Почта тумо", "Создан аккаунт с графиком и коучем", ""]
        def kid(name, sched, coach, room, email, note=""):
            return ["Мама Ученика", "87001234567", "mama@gmail.com", name, "87770001122", "01.02.2013",
                    sched, coach, None, room, email, "True", note]
        tabs = {"4 поток": [header,
                            ["Понедельник, Четверг : 10:30–12:30"] + [None] * 12,  # section title row
                            kid("КУМАРБЕКОВ ДИАР", "Понедельник, Четверг : 10:30–12:30", "Ayaulym", "3 этаж, 7 WR",
                                "Diar.Kumarbekov@tumo.world"),
                            kid("Салибек Айша", " Вторник, Пятница : 16:30–18:30 ", "Aruzhan T", "2 этаж, зона SL", ""),
                            kid("Ушедший Ученик", "Вторник, Пятница : 14:30–16:30", "Alen", "3 этаж, 9 WR", "",
                                note="лист ожидания ноябрь"),
                            kid("Красный Ученик", "Вторник, Пятница : 14:30–16:30", "Alen", "3 этаж, 9 WR", "")],
                "5 поток": [header, kid("Пятый Поток", "Вторник, Пятница : 14:30–16:30", "Danel", "", "")]}
        raw = xlsx(tabs, {("4 поток", 6): "FF0000"})
        stats = {}
        students, problems = sts.parse(sts.read_tables(raw, streams=("4",)), stats=stats)
        self.assertEqual(students, [
            ["Кумарбеков Диар", "Понедельник", "10:30-12:30", "Ayaulym", "", "Онбординг",
             "diar.kumarbekov@tumo.world", "3 этаж, 7 WR", "", []],
            ["Салибек Айша", "Вторник", "16:30-18:30", "Aruzhan Torebek", "", "Онбординг", "", "2 этаж, зона SL", "", []],
            ["Ушедший Ученик", "", "", "", "", "Онбординг", "", "", "waitlist", []],
            ["Красный Ученик", "", "", "", "", "Онбординг", "", "", "inactive", []],
        ])
        self.assertEqual(problems, [])
        self.assertEqual(stats, {"red": 1, "purple": 0, "notes_inactive": 0, "notes_waitlist": 1})
        flat = repr(students)
        for private in ("Мама", "8700", "8777", "gmail", "01.02.2013"):
            self.assertNotIn(private, flat)
        # 5 поток is included by default
        with_5 = sts.parse(sts.read_tables(raw))[0]
        self.assertEqual(with_5[-1][:6], ["Пятый Поток", "Вторник", "14:30-16:30", "Danel", "", "Онбординг"])

    def test_workshop_sheet_is_attached_by_full_name(self):
        import datetime
        tabs = {"ПН ЧТ 1030": [HEADER,
                               row("Аскерұлы Есбер", "Понедельник", "Gulnaz", "Gulnaz", email="askeruly.esber@tumo.world"),
                               row("Иван Петров Сергеевич", "Четверг", "Almas", "Almas"),
                               row("Красный Ученик", "Понедельник", "Almas", "Almas"),
                               row("Айгерим Сапарова", "Понедельник", "Almas", "Almas"),
                               row("Айгерим Сапарова Ерлановна", "Четверг", "Almas", "Almas")]}
        master = xlsx(tabs, {("ПН ЧТ 1030", 4): "FF0000"})
        workshops = xlsx({"Лист1": [
            ["ФИО ученика", "Воркшоп", "Кабинет / Зона", "Коуч", "Дни и время", "Дата начала", "Дата окончания",
             "Дата изменения", "Телефон"],
            ["АСКЕРУЛЫ ЕСБЕР", "Game Development 2", "3 этаж, 5 WR", "Ахмет", "Вт, Пт 16:30–18:30",
             datetime.datetime(2026, 11, 2), datetime.datetime(2026, 11, 27), "29.10.2026", "87001234567"],
            ["Аскерұлы Есбер", "Robotics 1", "2 этаж", "Dana Kurak", "Вт, Пт 16:30–18:30", "01.10.2026", "31.10.2026", "", ""],
            ["Иван Петров", "Music 1", "", "", "", "", "", "", ""],        # no patronymic, unique -> matched
            ["Айгерим Сапарова", "Film 1", "", "", "", "", "", "", ""],    # exact name wins over the longer one
            ["Красный Ученик", "Design 1", "", "", "", "", "", "", ""],    # inactive -> not attached
            ["Неизвестный Ученик", "3D 1", "", "", "", "", "", "", ""],
            [None, None, None, None, None, None, None, None, None]]})
        students, _ = sts.parse(sts.read_tables(master))
        rows = sts.read_workshop_rows(workshops)
        self.assertEqual(len(rows), 6)
        unmatched = sts.attach_workshops(students, rows)
        by_name = {s[0]: s for s in students}
        self.assertEqual(by_name["Аскерұлы Есбер"][9], [
            ["Robotics 1", "2 этаж", "Dana Kurak", "Вт, Пт 16:30–18:30", "2026-10-01", "2026-10-31", ""],
            ["Game Development 2", "3 этаж, 5 WR", "Akhmet", "Вт, Пт 16:30–18:30", "2026-11-02", "2026-11-27", "2026-10-29"],
        ])  # sorted by start date; coach stored like master-sheet coaches; the phone column is never read
        self.assertNotIn("87001234567", repr(students))
        self.assertEqual([w[0] for w in by_name["Иван Петров Сергеевич"][9]], ["Music 1"])
        self.assertEqual([w[0] for w in by_name["Айгерим Сапарова"][9]], ["Film 1"])
        self.assertEqual(by_name["Айгерим Сапарова Ерлановна"][9], [])
        self.assertEqual(by_name["Красный Ученик"][9], [])
        self.assertEqual([w["title"] for w in unmatched], ["Design 1", "3D 1"])
        # the master sheet's own «Воркшопы» tab is read too; no such tab / CSV -> nothing
        self.assertEqual(sts.read_workshop_rows(master, sts.WORKSHOP_TAB_RE), [])
        self.assertEqual(sts.read_workshop_rows(b"a,b\n1,2\n"), [])

    def test_workshops_survive_save_and_load(self):
        import tempfile
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "students.json"
            ws = [["Game Development 2", "3 этаж", "Ахмет", "Вт, Пт 16:30–18:30", "2026-11-02", "2026-11-27", ""]]
            sts.save_student_data([["Иван Петров", "Вторник", "10:30-12:30", "Almas", "", "Самообучение", "", "", "", ws]], path)
            self.assertEqual(sts.load_student_data(path)[0][9], ws)

    def test_html_login_page_is_rejected(self):
        with self.assertRaises(SystemExit):
            sts.read_tables(b"<!DOCTYPE html><html>Sign in</html>")


if __name__ == "__main__":
    unittest.main()
