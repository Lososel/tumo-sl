#!/usr/bin/env python3
"""Build data/students.json from the master student spreadsheet
("MAIN_ОБЩИЙ_1,2,3 БАТЧ_ВСЕХ_УЧЕНИКОВ_SL + ON").

The master sheet contains private data (parent/student phones, IIN, personal emails,
internal notes). Only the fields the public page needs are copied:
name, SL day, time slot, coach, ex-coach, stage and the student's @tumo.world email.

Input: the workbook as .xlsx (all schedule tabs) or a CSV export of a single tab.
  - xlsx: every tab named like "ПН ЧТ 1030" / "ВТ ПТ 1630" is read; the tab name gives the time slot.
    Onboarding stream tabs ("4 поток", see --streams) are read too: name from «ФИО ребенка»,
    days and slot from «Расписание», coach from «Коуч», room from «Зона обучения», stage «Онбординг».
  - csv:  the time slot is taken from "Новое расписание" or "Расписание".

Status of a student (row colour = fill of the «ФИО» cell; .xlsx only, a CSV has no colours):
  - RED row -> "inactive": the student left. The site shows only the name and «Если вы
    хотите продолжить обучение в TUMO, свяжитесь с нами через WhatsApp».
  - DARK PURPLE row (#9900FF, #674EA7, #351C75...) -> "waitlist": the account is suspended
    for now. The site shows only the name and «Сейчас ваш аккаунт приостановлен…».
    Light purples (#B4A7D6, #D9D2E9) are room groups in the "N поток" tabs, not a status.
  - Notes/status columns («Причина ухода», «Ушел к коучу», unlabelled note columns...):
    "отказ", "отчислен", "ignore", "self initiated", "не может/не сможет посещать" ->
    "inactive"; "приостановлено", "suspend", or a waitlist for a later period ("ЛИСТ
    ОЖИДАНИЯ 2027", "лист ожидания лето/январь/ноябрь", "на след./новый поток", "после НГ")
    -> "waitlist".
    Active: "лист ожидания", "лист ожидания (хочет пн/чт 16:30)", "хочет на 16:30"... — the
    student attends now and only wants another slot.
  Inactive/waitlist rows keep only the name (no schedule, coach or email).

Workshops (optional): a separate workshop sheet (--workshops, e.g. «Workshops 26.10») and/or
  a «Воркшопы…» tab in the master sheet. Columns (header row, any order):
    ФИО ученика | Почта TUMO (optional) | Воркшоп | Кабинет / Зона | Коуч | Дни и время |
    Дата начала | Дата окончания | Дата изменения
  One row per student per workshop (current and next). Matched to students by full name
  (TUMO email wins when given). The card shows «Текущий / Следующий воркшоп» and alerts for a
  new workshop starting soon or a recent change. No rows = no workshop on the cards.
  The «Расписание» column is not checked: "(лист ожидания)" there is the slot the family
  first asked for, not the student's current status.

Usage:
  python3 tools/sheet_to_students.py master.xlsx            # write data/students.json if changed
  python3 tools/sheet_to_students.py export.csv --dry-run   # only print the summary
Run automatically by .github/workflows/sync-students.yml.
"""
import argparse
import csv
import io
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from sync_schedule import (  # noqa: E402
    CYR_TO_LATIN_COACH, DAYS_OK, STUDENTS_JSON, clean_ws, load_student_data, norm, proper_case,
    save_student_data,
)

STAGE = "Самообучение"
ONBOARDING_STAGE = "Онбординг"
STREAM_TAB_RE = re.compile(r"^(\d+)\s*поток$", re.I)
DEFAULT_STREAMS = ("4", "5")  # onboarding stream tabs to include, by number
SLOT_BY_HOUR = {"10": "10:30-12:30", "14": "14:30-16:30", "16": "16:30-18:30"}
TAB_RE = re.compile(r"^(ПН\s*ЧТ|ВТ\s*ПТ)\s*(\d{2})\d{2}$", re.I)
DAYS_BY_PAIR = {"пнчт": {"Понедельник", "Четверг"}, "втпт": {"Вторник", "Пятница"}}
TUMO_EMAIL_RE = re.compile(r"^[a-z0-9._%+-]+@tumo\.world$")

# Notes/status columns checked for INACTIVE_RE (prefix match), plus columns with an empty header.
NOTE_HEADERS = ("причина ухода", "ушел к коучу", "ушёл к коучу", "статус", "примечан", "коммент", "status", "notes")
_LATER = r"(?:20[2-3]\d|лет|зим|весн|осен|январ|феврал|март|ноябр|декабр|след|нов\w*\s+поток)"
DROPPED_RE = re.compile(
    r"ignore|self[\s-]*initiated|отчисл|отказ|не\s+(?:может|сможет|будет)\s+посещать", re.I)
PAUSED_RE = re.compile(
    r"приостанов|suspend"
    r"|(?:лист\w*\s+ожидани\w*|waitlist|\bwl\b).*?" + _LATER +
    r"|(?:след\w*|нов\w*)\s+поток|после\s+нг",
    re.I,
)
INACTIVE_RE = re.compile(DROPPED_RE.pattern + "|" + PAUSED_RE.pattern, re.I)  # any non-active note
INACTIVE, WAITLIST = "inactive", "waitlist"
STATUS_RANK = {"": 0, WAITLIST: 1, INACTIVE: 2}  # a student in several rows: the most active wins


def note_status(text):
    if DROPPED_RE.search(text):
        return INACTIVE
    if PAUSED_RE.search(text):
        return WAITLIST
    return ""


def is_red(color):
    """openpyxl Color -> True for red highlights (#FF0000, #E06666, #CC0000...), not for orange,
    magenta, pink or yellow. Theme colours can't be resolved without the theme and count as not red."""
    if color is None:
        return False
    rgb = None
    if color.type == "rgb" and isinstance(color.rgb, str):
        rgb = color.rgb[-6:]
    elif color.type == "indexed" and color.indexed is not None:
        from openpyxl.styles.colors import COLOR_INDEX

        if 0 <= color.indexed < len(COLOR_INDEX):
            rgb = COLOR_INDEX[color.indexed][-6:]
    if not rgb or not re.fullmatch(r"[0-9A-Fa-f]{6}", rgb):
        return False
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    return r >= 200 and g <= 115 and b <= 115 and r - max(g, b) >= 90


def _rgb(color):
    if color is None:
        return None
    rgb = None
    if color.type == "rgb" and isinstance(color.rgb, str):
        rgb = color.rgb[-6:]
    elif color.type == "indexed" and color.indexed is not None:
        from openpyxl.styles.colors import COLOR_INDEX

        if 0 <= color.indexed < len(COLOR_INDEX):
            rgb = COLOR_INDEX[color.indexed][-6:]
    if not rgb or not re.fullmatch(r"[0-9A-Fa-f]{6}", rgb):
        return None
    return tuple(int(rgb[i:i + 2], 16) for i in (0, 2, 4))


def is_dark_purple(color):
    """#9900FF (Sheets "purple"), #674EA7 / #351C75 ("dark purple 1/2") -> True. Light purples
    (#B4A7D6, #D9D2E9: room groups), magenta (#FF00FF) and blues -> False."""
    import colorsys

    rgb = _rgb(color)
    if rgb is None:
        return False
    h, l, s = colorsys.rgb_to_hls(*(v / 255 for v in rgb))
    return 0.70 <= h <= 0.80 and s >= 0.35 and l <= 0.55


def _fill_colors(cell):
    fill = getattr(cell, "fill", None)
    if fill is None or fill.fill_type in (None, "none"):
        return ()
    return (fill.fgColor, fill.start_color)


def red_fill(cell):
    return any(is_red(c) for c in _fill_colors(cell))


def fill_status(cell):
    colors = _fill_colors(cell)
    if any(is_red(c) for c in colors):
        return INACTIVE
    if any(is_dark_purple(c) for c in colors):
        return WAITLIST
    return ""


# Column header (lowercased, prefix match) -> field. Anything else in the sheet is ignored.
HEADERS = {
    "фио ученика": "name",
    "фио ребенка": "name",
    "фио ребёнка": "name",
    "день sl": "day",
    "новый коуч sl": "coach",
    "коуч": "coach",
    "ex-coach": "ex",
    "почта tumo": "email",
    "почта тумо": "email",
    "зона обучения": "room",
    "новое расписание": "new_schedule",
    "расписание": "schedule",
}


# Latin letters that look like Cyrillic ones, typed by mistake inside Cyrillic names
# ("Tөлеуғазы" with a Latin T, "Iңкар" with a Latin I). Fixed only in words that mix both scripts,
# so Latin names ("Agayeva Zakhirat") stay as they are.
LATIN_TO_CYRILLIC = str.maketrans("ABCEHIKMOPTXYabceikmopxy", "АВСЕНІКМОРТХУавсеікморху")
CYRILLIC_RE = re.compile(r"[\u0400-\u04ff]")


def fix_mixed_script(name):
    return " ".join(w.translate(LATIN_TO_CYRILLIC) if CYRILLIC_RE.search(w) else w for w in name.split(" "))


def slot_from_text(text):
    m = re.search(r"(\d{1,2})[:.]\d{2}", str(text or ""))
    return SLOT_BY_HOUR.get(m.group(1).zfill(2)) if m else None


def canon_coach(raw):
    s = clean_ws(raw)
    return CYR_TO_LATIN_COACH.get(s.lower(), proper_case(s)) if s else ""


def map_header(row):
    cols = {}
    for i, cell in enumerate(row):
        h = clean_ws(cell).lower()
        for prefix, field in HEADERS.items():
            if h.startswith(prefix) and field not in cols:
                cols[field] = i
                break
    return cols


WORKSHOP_TAB_RE = re.compile(r"^\s*(воркшоп|workshop)", re.I)
# Header (lowercased) -> field; the longest matching prefix wins («дата начала» before «дата»).
WORKSHOP_HEADERS = {
    "фио": "name", "ученик": "name", "name": "name", "student": "name",
    "почта tumo": "email", "почта тумо": "email", "email": "email",
    "воркшоп": "workshop", "название": "workshop", "workshop": "workshop", "направление": "workshop", "тема": "workshop",
    "кабинет": "room", "зона": "room", "аудитория": "room", "room": "room",
    "коуч": "coach", "ведущий": "coach", "лидер": "coach", "тренер": "coach", "coach": "coach", "leader": "coach",
    "workshop leader": "coach", "workshop coach": "coach",
    "дни и время": "time", "расписание": "time", "время": "time", "дни": "time", "schedule": "time", "time": "time",
    "дата начала": "start", "начало": "start", "с даты": "start", "start": "start",
    "дата окончания": "end", "окончание": "end", "конец": "end", "по дату": "end", "end": "end",
    "дата изменения": "changed", "изменено": "changed", "изменен": "changed", "changed": "changed", "updated": "changed",
}
WORKSHOP_FIELDS = ["title", "room", "coach", "time", "start", "end", "changed"]  # one workshop in students.json


def _iso_date(value):
    """Cell value -> "YYYY-MM-DD" ("" if not a date). Accepts dates and dd.mm.yyyy / dd.mm text."""
    import datetime

    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.strftime("%Y-%m-%d")
    m = re.search(r"(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?", str(value or ""))
    if not m:
        return ""
    year = int(m.group(3)) if m.group(3) else datetime.date.today().year
    year += 2000 if year < 100 else 0
    try:
        return datetime.date(year, int(m.group(2)), int(m.group(1))).isoformat()
    except ValueError:
        return ""


def _workshop_columns(header):
    cols = {}
    for i, cell in enumerate(header):
        h = clean_ws(cell).lower()
        hits = [p for p in WORKSHOP_HEADERS if h.startswith(p)]
        if hits:
            field = WORKSHOP_HEADERS[max(hits, key=len)]
            cols.setdefault(field, i)
    return cols


def read_workshop_rows(raw, tab_re=None):
    """xlsx bytes -> [{name, email, title, room, coach, time, start, end, changed}].

    Reads every tab whose header has a «Воркшоп» column and a «ФИО ученика» (or «Почта TUMO»)
    column; tab_re limits the tabs (the master sheet: only «Воркшопы…» tabs). Other columns
    (phones, notes...) are ignored. [] for CSV input or when nothing matches."""
    if raw[:2] != b"PK":
        return []
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        if tab_re is not None and not tab_re.match(ws.title):
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        cols = _workshop_columns(rows[0])
        if "workshop" not in cols or not ({"name", "email"} & set(cols)):
            continue
        for line_no, r in enumerate(rows[1:], start=2):
            get = lambda f: r[cols[f]] if f in cols and cols[f] < len(r) else None  # noqa: E731
            title = clean_ws(get("workshop"))
            if not title:
                continue
            out.append({
                "where": f"[{ws.title}:{line_no}]",
                "name": proper_case(fix_mixed_script(clean_ws(get("name")))),
                "email": clean_ws(get("email")).lower(),
                "title": title, "room": clean_ws(get("room")),
                "coach": canon_coach(get("coach")), "time": clean_ws(get("time")),
                "start": _iso_date(get("start")), "end": _iso_date(get("end")), "changed": _iso_date(get("changed")),
            })
    return out


def attach_workshops(students, workshop_rows):
    """Add each workshop row to its student's workshop list (field 9), sorted by start date.

    The full name is the key (Kazakh/Russian letters, case and word order don't matter); a
    TUMO email, when given, wins; a name without the patronymic matches if only one student
    fits. Inactive / waitlist students get nothing. Returns the rows that matched no student."""
    active = [s for s in students if not s[8]]
    by_email = {s[6]: s for s in active if s[6]}
    by_name = {}
    for s in active:
        by_name.setdefault(norm(s[0]), []).append(s)
    tokens = [(frozenset(norm(s[0]).split()), s) for s in active]
    unmatched = []
    for w in workshop_rows:
        s = by_email.get(w["email"]) if w["email"] else None
        if s is None:
            same = by_name.get(norm(w["name"]), [])
            if len(same) == 1:
                s = same[0]
            elif not same:
                want = frozenset(norm(w["name"]).split())
                fits = [st for tk, st in tokens if len(want) >= 2 and want <= tk]
                s = fits[0] if len(fits) == 1 else None
        if s is None:
            unmatched.append(w)
            continue
        entry = [w[f] for f in WORKSHOP_FIELDS]
        if entry not in s[9]:
            s[9].append(entry)
            s[9].sort(key=lambda e: (e[4] or "9999", e[0]))
    return unmatched


def stream_of(tab):
    m = STREAM_TAB_RE.match(tab or "")
    return m.group(1) if m else None


def read_tables(raw, streams=DEFAULT_STREAMS):
    """bytes -> [(tab_name or None, rows, colors)] where colors[i] is the status the fill of
    row i's «ФИО» cell stands for: "inactive" (red), "waitlist" (dark purple) or "". CSV: None."""
    if raw[:2] == b"PK":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        tables = []
        for ws in wb.worksheets:
            title = ws.title.strip()
            if not (TAB_RE.match(title) or stream_of(title) in set(streams)):
                continue
            rows, red, name_col = [], [], 0
            for i, cells in enumerate(ws.iter_rows()):
                values = tuple(c.value for c in cells)
                if i == 0:
                    name_col = map_header(values).get("name", 0)
                rows.append(values)
                red.append(fill_status(cells[name_col]) if i > 0 and name_col < len(cells) else "")
            tables.append((ws.title.strip(), rows, red))
        return tables
    text = raw.decode("utf-8-sig")
    if text.lstrip()[:1] == "<":
        raise SystemExit("Got an HTML page instead of the sheet: check the link/sharing or the credentials.")
    return [(None, list(csv.reader(io.StringIO(text))), None)]


def note_columns(header, name_col):
    """Indexes of notes/status columns: known headers, or unlabelled columns after the name."""
    cols = []
    for i, cell in enumerate(header):
        h = clean_ws(cell).lower()
        if (h and h.startswith(NOTE_HEADERS)) or (not h and i > name_col):
            cols.append(i)
    return cols


def pick_duplicate(cands, published):
    """A student listed in several tabs (usually moved between slots). Prefer, in order:
    an active row over a waitlist row over an inactive one, the row whose "Новое расписание"
    confirms its tab slot, the row matching what is published now (so an ambiguous sheet
    doesn't flip the schedule), then the last row."""
    best = min(STATUS_RANK[c["rec"][8]] for c in cands)
    cands = [c for c in cands if STATUS_RANK[c["rec"][8]] == best]
    if len(cands) == 1:
        return cands[0]
    confirmed = [c for c in cands if c["confirmed"]]
    if len(confirmed) == 1:
        return confirmed[0]
    if published:
        same = [c for c in cands if c["rec"][1:3] == published[1:3]]
        if same:
            return same[-1]
    return cands[-1]


def parse(tables, published=None, stats=None):
    """-> (students, problems). published: {norm(name): row} of the current data/students.json.
    stats (optional dict) receives the counts of inactive / waitlist rows by reason."""
    published = published or {}
    stats = stats if stats is not None else {}
    stats.update(red=0, purple=0, notes_inactive=0, notes_waitlist=0)
    by_key, problems = {}, []
    for table in tables:
        tab, rows = table[0], table[1]
        colors = table[2] if len(table) > 2 and table[2] else [""] * len(rows)
        if not rows:
            continue
        cols = map_header(rows[0])
        onboarding = stream_of(tab) is not None
        if "name" not in cols or ("schedule" if onboarding else "day") not in cols:
            problems.append(f"[{tab}] нет колонок «ФИО» / «{'Расписание' if onboarding else 'День SL'}» — вкладка пропущена")
            continue
        notes = note_columns(rows[0], cols["name"])
        m = TAB_RE.match(tab or "")
        tab_slot = SLOT_BY_HOUR.get(m.group(2)) if m else None
        pair_days = DAYS_BY_PAIR[re.sub(r"\s", "", m.group(1)).lower()] if m else set(DAYS_OK)

        def get(row, field):
            i = cols.get(field)
            return row[i] if i is not None and i < len(row) else None

        for line_no, row in enumerate(rows[1:], start=2):
            name = proper_case(fix_mixed_script(clean_ws(get(row, "name"))))
            if not name:
                continue
            where = f"[{tab or 'csv'}:{line_no}] {name}"
            stage = ONBOARDING_STAGE if onboarding else STAGE
            status = colors[line_no - 1] or ""
            if status:
                stats["red" if status == INACTIVE else "purple"] += 1
            else:
                for i in notes:
                    if i < len(row) and row[i]:
                        status = note_status(str(row[i]))
                        if status:
                            stats["notes_" + status] += 1
                            break
            if status:
                # Only the name: the card says «свяжитесь с нами», no schedule/coach/email.
                rec = [name, "", "", "", "", stage, "", "", status, []]
                by_key.setdefault(norm(name), []).append({"rec": rec, "where": where, "confirmed": False})
                continue
            if onboarding:
                # «Понедельник, Четверг : 10:30–12:30»: the pair's first day stands for the pair
                # (onboarding has no main SL day yet); the slot comes from the start time.
                sched = clean_ws(get(row, "schedule")).lower()
                day = "Понедельник" if "понедельник" in sched or "четверг" in sched else (
                      "Вторник" if "вторник" in sched or "пятница" in sched else sched.capitalize())
                pair_days = set(DAYS_OK)
            else:
                day = clean_ws(get(row, "day")).capitalize()
            time = tab_slot or slot_from_text(get(row, "new_schedule")) or slot_from_text(get(row, "schedule"))
            if day not in DAYS_OK or day not in pair_days:
                problems.append(f"{where}: день SL не распознан: {day!r}")
                continue
            if not time:
                problems.append(f"{where}: время не распознано")
                continue
            email = clean_ws(get(row, "email")).lower()
            rec = [name, day, time, canon_coach(get(row, "coach")), canon_coach(get(row, "ex")), stage,
                   email if TUMO_EMAIL_RE.match(email) else "", clean_ws(get(row, "room")), "", []]
            confirmed = bool(tab_slot) and slot_from_text(get(row, "new_schedule")) == tab_slot
            by_key.setdefault(norm(name), []).append({"rec": rec, "where": where, "confirmed": confirmed})

    students = []
    for key, cands in by_key.items():
        if len({tuple(c["rec"][:9]) for c in cands}) > 1:
            chosen = pick_duplicate(cands, published.get(key))
            # A red/purple old row next to the current one is expected; warn only when rows
            # of the same status disagree, i.e. the sheet is really ambiguous.
            same = {tuple(c["rec"][:9]) for c in cands if c["rec"][8] == chosen["rec"][8]}
            if len(same) == 1:
                students.append(chosen["rec"])
                continue
            problems.append(f"{chosen['where']}: ученик есть в нескольких строках ("
                            + "; ".join(c["where"].split("]")[0] + "]" for c in cands)
                            + f") — выбрано {chosen['rec'][1]} {chosen['rec'][2]}")
            students.append(chosen["rec"])
        else:
            students.append(cands[0]["rec"])
    return students, problems


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", help="Downloaded master sheet (.xlsx or .csv)")
    ap.add_argument("--out", default=str(STUDENTS_JSON))
    ap.add_argument("--min-students", type=int, default=200,
                    help="Refuse to write if fewer students were parsed (protects against a broken download)")
    ap.add_argument("--max-drop", type=float, default=0.2,
                    help="Refuse to write if the student count falls by more than this share "
                         "(e.g. a single-tab CSV instead of the whole workbook)")
    ap.add_argument("--streams", default=",".join(DEFAULT_STREAMS),
                    help='Onboarding stream tabs to include, e.g. "4" or "4,5" (default: %(default)s)')
    ap.add_argument("--workshops", action="append", default=[], metavar="FILE",
                    help="Downloaded workshop sheet (.xlsx); may be given several times")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    out = Path(args.out)
    current = load_student_data(out) if out.exists() else []
    stats = {}
    streams = tuple(s.strip() for s in args.streams.split(",") if s.strip())
    raw = Path(args.source).read_bytes()
    students, problems = parse(read_tables(raw, streams), {norm(r[0]): r for r in current}, stats)
    workshop_rows = read_workshop_rows(raw, WORKSHOP_TAB_RE)
    for extra in args.workshops:
        if Path(extra).exists():
            workshop_rows += read_workshop_rows(Path(extra).read_bytes())
    unmatched = attach_workshops(students, workshop_rows)
    if workshop_rows:
        print(f"Воркшопы: {len(workshop_rows)} строк, найдено учеников: {len(workshop_rows) - len(unmatched)}")
        for w in unmatched:
            problems.append(f"{w['where']} воркшоп «{w['title']}»: ученик не найден: {w['email'] or w['name']}")
    active = [s for s in students if not s[8]]
    onboarding = sum(1 for s in active if s[5] == ONBOARDING_STAGE)
    print(f"Учеников: {len(students)} — активных: {len(active)} (самообучение: {len(active) - onboarding}, "
          f"онбординг: {onboarding}), неактивных: {sum(1 for s in students if s[8] == INACTIVE)}, "
          f"в листе ожидания: {sum(1 for s in students if s[8] == WAITLIST)}; проблем: {len(problems)}")
    print(f"Строки по статусу — неактивные: красная заливка {stats['red']}, отметка в примечаниях "
          f"{stats['notes_inactive']}; лист ожидания: фиолетовая заливка {stats['purple']}, отметка в "
          f"примечаниях {stats['notes_waitlist']}")
    for p in problems:
        print("  " + p)
    if len(students) < args.min_students:
        raise SystemExit(f"Слишком мало учеников ({len(students)} < {args.min_students}) — файл не записан.")

    if current and len(students) < len(current) * (1 - args.max_drop):
        raise SystemExit(f"Учеников стало {len(students)} вместо {len(current)} (> {args.max_drop:.0%} меньше) — "
                         "файл не записан. Проверьте, что выгружается вся книга, а не одна вкладка, "
                         "или запустите с --max-drop 1.")
    if current == students:
        print("Без изменений.")
        return
    if args.dry_run:
        print("Dry-run: файл не записан.")
        return
    save_student_data(students, out)
    print(f"Записано: {out}")


if __name__ == "__main__":
    main()
